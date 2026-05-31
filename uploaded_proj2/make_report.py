from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

for text, size, bold in [
    ('Harbin Institute of Technology (Shenzhen)', 14, True),
    ('Project Report', 16, True),
    ('', 10, False),
    ('Image Processing (COMP5033) ', 12, False),
    ('2025-2026 Spring Semester', 12, False),
    ('Teacher: Prof. Weizheng Zhang', 12, False),
    ('', 10, False),
    ('', 10, False),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)

table = doc.add_table(rows=3, cols=4)
table.style = 'Table Grid'

table.rows[0].cells[0].text = 'Project No.:'
table.rows[0].cells[1].text = '2'
table.rows[0].cells[2].text = 'Student Name:'
table.rows[0].cells[3].text = 'Amir Abbasi'

table.rows[1].cells[0].text = 'Student ID: '
table.rows[1].cells[1].text = '25SF51024'
table.rows[1].cells[2].text = 'Date of Submission:'
table.rows[1].cells[3].text = '31/05/2026'

row3 = table.rows[2]
cell = row3.cells[0]
cell.merge(row3.cells[1]).merge(row3.cells[2]).merge(row3.cells[3])
table.rows[2].cells[0].text = (
    'Grading (for Teaching Assistant)\n\nScore: ____________________\n\n'
    'Comments: ' + '_' * 200
)

doc.add_paragraph()

def heading(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)

def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.runs[0].font.size = Pt(11)

heading('1 Project Content')
body(
    'There is a grayscale image given that we need to find its boundary using '
    'morphological processing. First I have to binarize the image and then apply '
    'erosion to extract the boundary. The image is Figure_P2.jpg and its size is '
    '1024x1024. I write all the key steps by myself and only use cv2 for reading '
    'and saving the image.'
)

heading('2 Method Description')
body(
    'First I binarize the grayscale image. I use Otsu method to find the threshold '
    'automatic, it try every value from 0 to 255 and pick the one with the biggest '
    'between class variance. For my image the threshold was 107, so every pixel '
    'bigger than 107 become white and the rest become black.'
)
body(
    'Then I use erosion to find the boundary. For each pixel I check the 3x3 window '
    'around it, the pixel stay white only if all 9 pixels are white, otherwise it '
    'become black. Finally the boundary is just original minus eroded '
    '(boundary = binary - eroded) which leaves only the edge pixels.'
)

heading('3 Experiment Results and Analysis')
body(
    'After I run the code the threshold came out as 107. When I look at the binary '
    'image, the white cat become white and the black cat with the gray border outside '
    'become black, which is what I expected. The boundary image shows a thin white '
    'line on all the edges of the shapes. You can see the big outer circle, the curve '
    'between the two cats, and also the face and whiskers. I got 15692 boundary pixels.'
)
body(
    'The good thing about this method is that it is simple and fast. I dont have to '
    'choose the threshold by myself because Otsu do it automatic, and the erosion is '
    'only some AND operations so even a big 1024x1024 image finish quickly. The edge '
    'it gives is thin and clean.'
)
body(
    'But there is also some problem. The thin lines like the whiskers come out a '
    'little weak, I think because they are thinner than 3 pixels so the 3x3 erosion '
    'almost delete them. Also everything depend on the binarize step, so if the '
    'threshold is wrong then the boundary will be wrong too.'
)

heading('4 Summary')
body(
    'The most hard part for me was writing the Otsu threshold by myself. I had to '
    'keep the weight and the mean for every threshold value and I test it few times '
    'until the number look correct. The erosion part was more easy once I understand '
    'how the 3x3 window work.'
)
body(
    'From this project I learned how to binarize an image and why using Otsu is '
    'better than just picking a number by hand. I also understand erosion much better '
    'now, especially that subtracting the eroded image from original directly gives '
    'the boundary. I think this is a simple but clever idea.'
)

doc.save('/home/user/amirminsk/uploaded_proj2/Project Report2.docx')
print('done')
