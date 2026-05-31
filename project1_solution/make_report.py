from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

# title
for text, size, bold, center in [
    ('Harbin Institute of Technology (Shenzhen)', 14, True, True),
    ('Project Report', 16, True, True),
    ('Image Processing (COMP5033)', 12, False, True),
    ('2025–2026 Spring Semester', 12, False, True),
    ('Lecturer: Prof. Weizheng Zhang', 12, False, True),
]:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)

doc.add_paragraph()

# info table
table = doc.add_table(rows=2, cols=4)
table.style = 'Table Grid'
table.rows[0].cells[0].text = 'Project No.:'
table.rows[0].cells[1].text = '1'
table.rows[0].cells[2].text = 'Student Name:'
table.rows[0].cells[3].text = '[YOUR NAME]'
table.rows[1].cells[0].text = 'Student ID:'
table.rows[1].cells[1].text = '[YOUR ID]'
table.rows[1].cells[2].text = 'Date of Submission:'
table.rows[1].cells[3].text = '2026-05-30'

doc.add_paragraph()

def heading(text, size=13):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)

def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    run = p.runs[0]
    run.font.size = Pt(11)

def add_img(path, caption, width=6):
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(10)

heading('1  Project Content')
body(
    'In this project we are given three noisy images and we need to figure out what '
    'type of noise is in each one and then remove it. We can\'t use the built-in '
    'processing functions from cv2, so I implemented everything myself using loops '
    'and basic math.'
)

heading('2  Method Description')

heading('Image 1 – Median Filter', size=11)
body(
    'Looking at image 1 I can clearly see white dots randomly scattered across the image. '
    'This is salt noise where some pixels get replaced with a very bright value.'
)
body(
    'I used a median filter with a 5×5 window. For each pixel I collect all 25 values '
    'in the surrounding area, sort them, and take the middle one. I first tried 3×3 '
    'but some noise clusters were still visible, so I increased it to 5×5 and that '
    'worked better. The median is good here because one or two bright outliers '
    'don\'t affect the result much.'
)

heading('Image 2 – Gaussian Filter', size=11)
body(
    'Image 2 has fine random noise all over it, not isolated dots like image 1. '
    'The whole image looks grainy. This is Gaussian noise which comes from the camera sensor.'
)
body(
    'I built a 5×5 Gaussian kernel using the formula: w(x,y) = exp(-(x²+y²) / (2σ²)) '
    'then normalized it so all weights add up to 1. I used σ=1.5. Then I applied it '
    'by doing a weighted sum over each pixel\'s neighbourhood. The result is a bit '
    'blurry but the noise is gone.'
)

heading('Image 3 – Frequency Domain Filter', size=11)
body(
    'Image 3 has a repeating dot pattern which looked like periodic noise to me. '
    'I learned in class that periodic noise shows up as bright spots in the frequency '
    'spectrum, so the best way to remove it is to filter in the frequency domain.'
)
body(
    'I took the 2D FFT of the image, shifted it so the low frequencies are in the '
    'centre, then applied a Gaussian low-pass mask. The mask reduces the high frequencies '
    'where the noise lives. Then I did the inverse FFT to get back the filtered image.'
)

heading('3  Experiment Results and Analysis')

add_img('fig_image1.png', 'Figure 1. Image 1 before and after median filter (5×5).')
doc.add_paragraph()
add_img('fig_image2.png', 'Figure 2. Image 2 before and after Gaussian filter (5×5, σ=1.5).')
doc.add_paragraph()
add_img('fig_image3.png', 'Figure 3. Image 3 before and after frequency domain low-pass filter.')
doc.add_paragraph()

body(
    'Image 1: the white dots are completely removed and the van and background look clean. '
    'Edges are still sharp which is what I expected from median filtering.'
)
body(
    'Image 2: the noise is reduced but the image is a bit soft. This is the tradeoff '
    'with Gaussian filtering – it smooths out noise but also blurs edges slightly. '
    'Using a smaller kernel would keep more detail but also more noise.'
)
body(
    'Image 3: the repeating dot pattern is mostly gone and the text is readable. '
    'I experimented with different cutoff values and 40 gave the best balance between '
    'removing the noise and keeping the image sharp.'
)

heading('4  Summary')
body(
    'I found the frequency domain filter the most interesting part. It was helpful to '
    'see how different types of noise look in the Fourier spectrum. The main difficulty '
    'was writing the median filter loop efficiently enough since 224×224 with a 5×5 '
    'window is a lot of iterations.'
)
body(
    'The key thing I learned is that you need to choose the right filter for the noise '
    'type – median for salt noise, Gaussian for random noise, and frequency domain for '
    'periodic patterns.'
)

doc.save('report_project1.docx')
print('saved report_project1.docx')
